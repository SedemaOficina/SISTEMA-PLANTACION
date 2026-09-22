"""Convierte el catálogo de especies (Excel del SIA) en assets/catalogo-especies.js.

El archivo fuente vive en assets/fuentes/ y no se edita: para cambiar una especie se corrige el
Excel y se vuelve a correr este script. Cada fila queda con el esquema del almacén `catalogos`
(tipo = especie): `id` y `clave` son el id_especie del catálogo (ESP-0000), que es la única
llave por la que se enlazan las plantaciones (D84); `nombre` es el nombre_comun, la etiqueta
de uso en campo. Los demás campos del catálogo se conservan con su nombre original.

Uso:  python3 pruebas/generar_especies.py
"""
import json, os, sys
import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FUENTE = os.path.join(RAIZ, 'assets', 'fuentes', 'CGO_ESPECIES_REFORESTACION_URBANA_2026-09-22.xlsx')
SALIDA = os.path.join(RAIZ, 'assets', 'catalogo-especies.js')
FECHA_CORTE = '2026-09-22'          # fecha de verificación contra EncicloVida, según el diccionario
VERSION = '2026-09-22'              # se sube cada vez que cambia el Excel

CAMPOS = ['id_especie', 'genero', 'especie', 'nombre_cientifico', 'nombre_comun', 'otros_nombres_comunes',
          'tipo_distribucion', 'id_snib', 'formadecrecimiento', 'id_enciclovida', 'nota_discrepancia']
DISTRIBUCION = {'Endémica', 'Nativa', 'Exótica', 'Exótica-Invasora'}

def limpiar(v):
    if v is None: return None
    if isinstance(v, str):
        v = ' '.join(v.split())
        return v or None
    return v

def main():
    wb = openpyxl.load_workbook(FUENTE, data_only=True)
    ws = wb['especies']
    hdr = [c.value for c in ws[1]]
    if hdr != CAMPOS:
        sys.exit('Encabezados distintos de los esperados:\n  %s\n  %s' % (hdr, CAMPOS))
    filas = [dict(zip(hdr, (limpiar(v) for v in r))) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

    # Validaciones: las que el diccionario de datos declara
    ids = [f['id_especie'] for f in filas]
    assert len(ids) == len(set(ids)), 'id_especie repetido'
    for f in filas:
        assert f['id_especie'].startswith('ESP-') and len(f['id_especie']) == 8, f['id_especie']
        for k in ('genero', 'especie', 'nombre_cientifico', 'nombre_comun', 'tipo_distribucion'):
            assert f[k], '%s sin %s' % (f['id_especie'], k)
        assert f['tipo_distribucion'] in DISTRIBUCION, (f['id_especie'], f['tipo_distribucion'])
        assert f['nombre_cientifico'].startswith(f['genero'] + ' '), (f['id_especie'], 'género no coincide')
        if f['id_enciclovida'] is not None: assert isinstance(f['id_enciclovida'], int), f['id_especie']
    comunes = [f['nombre_comun'] for f in filas]
    assert len(comunes) == len(set(comunes)), 'nombre_comun repetido'
    cientificos = [f['nombre_cientifico'] for f in filas]
    assert len(cientificos) == len(set(cientificos)), 'nombre_cientifico repetido'

    especies = []
    for f in filas:
        especies.append({
            'id': f['id_especie'], 'tipo': 'especie', 'clave': f['id_especie'],
            'nombre': f['nombre_comun'],
            'nombre_cientifico': f['nombre_cientifico'],
            'genero': f['genero'], 'especie': f['especie'],
            'otros_nombres_comunes': f['otros_nombres_comunes'] or '',
            'tipo_distribucion': f['tipo_distribucion'],
            'formadecrecimiento': f['formadecrecimiento'] or '',
            'id_snib': f['id_snib'] or None,
            'id_enciclovida': f['id_enciclovida'],
            'nota_discrepancia': f['nota_discrepancia'] or '',
            'activo': True, 'es_ficticio': False,
            'creado_por_id': None, 'fecha_creacion': FECHA_CORTE + 'T00:00:00-06:00',
            'editado_por_id': None, 'fecha_ultima_edicion': None,
        })

    meta = {'fuente': os.path.basename(FUENTE), 'fecha_corte': FECHA_CORTE, 'version': VERSION,
            'total': len(especies), 'verificado_contra': 'EncicloVida (CONABIO)',
            'siguiente_clave': 'ESP-%04d' % (max(int(i[4:]) for i in ids) + 1)}
    cuerpo = json.dumps({'meta': meta, 'especies': especies}, ensure_ascii=False, indent=1)
    with open(SALIDA, 'w', encoding='utf-8') as s:
        s.write('/* CATÁLOGO DE ESPECIES. Generado por pruebas/generar_especies.py a partir de\n'
                '   assets/fuentes/%s: no se edita a mano (D84). */\n' % os.path.basename(FUENTE))
        s.write('window.SRP = window.SRP || {};\nSRP.CATALOGO_ESPECIES = ' + cuerpo + ';\n')
    print('%d especies → %s (siguiente clave %s)' % (len(especies), os.path.relpath(SALIDA, RAIZ), meta['siguiente_clave']))

if __name__ == '__main__':
    main()
